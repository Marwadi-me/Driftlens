import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation with 16:9 widescreen dimensions
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette (Dark Mode Tech Theme)
BG_COLOR = RGBColor(9, 13, 22)        # Deep space dark #090D16
CARD_BG = RGBColor(18, 24, 38)        # Card Navy #121826
CARD_BORDER = RGBColor(31, 41, 61)    # Border Slate #1F293D
ACCENT_BLUE = RGBColor(56, 189, 248)  # Cyan #38BDF8
ACCENT_GREEN = RGBColor(16, 185, 129) # Emerald #10B981
ACCENT_ROSE = RGBColor(244, 63, 94)   # Rose Red #F43F5E
TEXT_MAIN = RGBColor(248, 250, 252)   # Near white #F8FAFC
TEXT_MUTED = RGBColor(148, 163, 184)  # Muted slate #94A3B8
HEADER_TAG = RGBColor(125, 211, 252)  # Light cyan #7DD3FC

blank_slide_layout = prs.slide_layouts[6]

def create_base_slide(title_text="", category_text="DRIFTLENS | HACKTOBERFEST 2026"):
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Background fill
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    
    if title_text:
        # Category / Header tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        # Slide Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.8))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = "Arial"
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

    return slide

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

# ==========================================
# SLIDE 1: Title Slide
# ==========================================
slide1 = prs.slides.add_slide(blank_slide_layout)
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg1.fill.solid()
bg1.fill.fore_color.rgb = BG_COLOR
bg1.line.fill.background()

# Badge Pill
pill = add_card(slide1, 0.8, 1.2, 5.2, 0.45, CARD_BG, ACCENT_BLUE)
pt = pill.text_frame.paragraphs[0]
pt.text = "★ Hacktoberfest 2026 Hack Day Sonipat — Problem 5"
pt.font.name = "Arial"
pt.font.size = Pt(12)
pt.font.bold = True
pt.font.color.rgb = ACCENT_BLUE
pt.alignment = PP_ALIGN.CENTER

# Title
t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(1.5))
tf = t_box.text_frame
p1 = tf.paragraphs[0]
p1.text = "DriftLens 🔍"
p1.font.name = "Arial"
p1.font.size = Pt(48)
p1.font.bold = True
p1.font.color.rgb = TEXT_MAIN

p2 = tf.add_paragraph()
p2.text = "Automated Documentation Drift Detection for Pull Requests"
p2.font.name = "Arial"
p2.font.size = Pt(24)
p2.font.color.rgb = ACCENT_BLUE
p2.space_before = Pt(8)

p3 = tf.add_paragraph()
p3.text = '"Catch the documentation a code change just made wrong."'
p3.font.name = "Arial"
p3.font.size = Pt(18)
p3.font.italic = True
p3.font.color.rgb = TEXT_MUTED
p3.space_before = Pt(12)

# Info Card 1
c1 = add_card(slide1, 0.8, 4.5, 3.6, 2.2)
c1_tf = c1.text_frame
c1_tf.word_wrap = True
p = c1_tf.paragraphs[0]
p.text = "Team LoneCommiter"
p.font.name = "Arial"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = c1_tf.add_paragraph()
p.text = "• Yashvardhan Mundhra\n• Role: Full-Stack Developer\n• Track: Developer Experience\n• Category: Static Analysis"
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)

# Info Card 2
c2 = add_card(slide1, 4.8, 4.5, 3.6, 2.2)
c2_tf = c2.text_frame
c2_tf.word_wrap = True
p = c2_tf.paragraphs[0]
p.text = "Open-Weight AI Model"
p.font.name = "Arial"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = c2_tf.add_paragraph()
p.text = "• Qwen 2.5 7B Instruct\n• 100% Hackathon Rule Compliant\n• Zero proprietary closed APIs\n• OpenAI-compatible vLLM/Ollama"
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)

# Info Card 3
c3 = add_card(slide1, 8.8, 4.5, 3.6, 2.2)
c3_tf = c3.text_frame
c3_tf.word_wrap = True
p = c3_tf.paragraphs[0]
p.text = "Key Deliverables"
p.font.name = "Arial"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = RGBColor(245, 158, 11)
p = c3_tf.add_paragraph()
p.text = "• CLI & GitHub Action Workflow\n• Interactive Dark-Mode Web Studio\n• Visual Blast Radius Impact Graph\n• 70 Unit & Integration Tests (100%)"
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)


# ==========================================
# SLIDE 2: The Core Problem
# ==========================================
s2 = create_base_slide("The Problem: Code Changes Invalidate Documentation")

prob_cards = [
    ("Flag Renames", "A CLI flag is renamed (--timeout -> --deadline), but prose guides and tutorials still instruct users to run the old command.", ACCENT_ROSE),
    ("Altered Defaults", "A default timeout changes from 30s to 60s in code, but configuration tables continue displaying the old 30s value.", ACCENT_ROSE),
    ("Signature Shifts", "API parameters are added or reordered (registerUser), yet integration tutorials continue explaining the legacy signature.", ACCENT_ROSE),
    ("Dead Features", "An obsolete endpoint (/api/v1/legacy-export) is deleted in code, but stays documented as active in API docs.", ACCENT_ROSE)
]

for idx, (head, desc, col) in enumerate(prob_cards):
    row = idx // 2
    col_idx = idx % 2
    x = 0.8 + col_idx * 5.9
    y = 1.8 + row * 2.5
    card = add_card(s2, x, y, 5.7, 2.2)
    ctf = card.text_frame
    ctf.word_wrap = True
    p = ctf.paragraphs[0]
    p.text = f"⚠️ {head}"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = col
    p = ctf.add_paragraph()
    p.text = desc
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(10)


# ==========================================
# SLIDE 3: Why Existing Tools Fail
# ==========================================
s3 = create_base_slide("Why Existing Tools Cannot Solve This Problem")

table_shape = s3.shapes.add_table(5, 4, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
table = table_shape.table
table.columns[0].width = Inches(2.2)
table.columns[1].width = Inches(2.5)
table.columns[2].width = Inches(4.0)
table.columns[3].width = Inches(3.0)

headers = ["Tool / Technology", "Primary Purpose", "How It Operates", "Fundamental Blindspot"]
for idx, h in enumerate(headers):
    cell = table.cell(0, idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = CARD_BORDER
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = ACCENT_BLUE

rows_data = [
    ("Python doctest", "Docstring test runner", "Executes isolated Python code snippets inside docstrings.", "Cannot evaluate markdown prose, tables, flags, or tutorials."),
    ("rustdoc tests", "Rust example testing", "Compiles and executes ```rust code blocks in documentation.", "Strictly limited to executable Rust code. Blind to prose statements."),
    ("mktestdocs / Doc Detective", "Markdown snippet runners", "Validates executable command fences and UI flows against live apps.", "Requires live servers; completely incapable of reasoning about PR diff contradictions."),
    ("DriftLens (Our Solution)", "Semantic Drift Detection", "AST entity extraction + Markdown chunking + Qwen 2.5 7B reasoning.", "Directly connects code diffs to prose sentences before merging.")
]

for row_idx, r in enumerate(rows_data, start=1):
    for col_idx, text in enumerate(r):
        cell = table.cell(row_idx, col_idx)
        cell.fill.solid()
        if row_idx == 4:
            cell.fill.fore_color.rgb = RGBColor(16, 42, 60)
        else:
            cell.fill.fore_color.rgb = CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(12)
        p.font.color.rgb = ACCENT_GREEN if row_idx == 4 else TEXT_MAIN
        if col_idx == 0:
            p.font.bold = True


# ==========================================
# SLIDE 4: The DriftLens Solution & Architecture
# ==========================================
s4 = create_base_slide("The DriftLens Architecture: 6-Stage Analysis Engine")

stages = [
    ("1. Git Diff Extraction", "Extracts unified diffs (base...head) via git/GitHub API, classifying files into source, documentation, and configuration.", ACCENT_BLUE),
    ("2. AST Entity Extraction", "Identifies modified constants, CLI flags (--timeout -> --deadline), function signatures, and deleted routes.", ACCENT_BLUE),
    ("3. Markdown Chunking", "Deconstructs Markdown files into sentence-level units with exact line ranges, sections, and table rows.", ACCENT_BLUE),
    ("4. Candidate Matching", "Multi-signal scoring: exact identifier (+45%), old value (+30%), section context (+15%), file correlation (+15%).", ACCENT_GREEN),
    ("5. Same-PR Suppression", "Checks if documentation was updated in the same PR. If already corrected, suppresses drift (0 false alarms).", ACCENT_GREEN),
    ("6. Qwen 2.5 7B Verification", "Open-weight model evaluates semantic contradiction with composite confidence calibration and 1-click suggested fix.", ACCENT_BLUE)
]

for idx, (title, desc, col) in enumerate(stages):
    row = idx // 3
    col_idx = idx % 3
    x = 0.8 + col_idx * 3.9
    y = 1.8 + row * 2.5
    card = add_card(s4, x, y, 3.8, 2.3)
    ctf = card.text_frame
    ctf.word_wrap = True
    p = ctf.paragraphs[0]
    p.text = title
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = col
    p = ctf.add_paragraph()
    p.text = desc
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(8)


# ==========================================
# SLIDE 5: Confidence Formula & Precision Engineering
# ==========================================
s5 = create_base_slide("Mathematical Confidence Calibration & False-Positive Shield")

# Left Column: The Formula
c_left = add_card(s5, 0.8, 1.8, 5.7, 4.8)
ltf = c_left.text_frame
ltf.word_wrap = True
p = ltf.paragraphs[0]
p.text = "Calibrated Composite Confidence Formula"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = ltf.add_paragraph()
p.text = "\nConfidence = 0.35 × S_det + 0.25 × S_cand + 0.40 × S_llm\n"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

bullets_left = [
    ("Deterministic Score (35%)", "1.0 for hard AST shifts (renamed flags, constant value changes); 0.80 for general changes."),
    ("Candidate Overlap Score (25%)", "Exact identifier match, old value mention in text, section relevance, file proximity."),
    ("LLM Semantic Score (40%)", "Open-weight Qwen 2.5 7B structured reasoning over tightly bounded evidence."),
    ("Strict Noise Suppression (< 0.60)", "Any score below 60% confidence is discarded to eliminate false alarm noise.")
]
for title, desc in bullets_left:
    p = ltf.add_paragraph()
    p.text = f"• {title}: {desc}"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(6)

# Right Column: Severity Tiers & Zero False Alarm Policy
c_right = add_card(s5, 6.8, 1.8, 5.7, 4.8)
rtf = c_right.text_frame
rtf.word_wrap = True
p = rtf.paragraphs[0]
p.text = "Actionable Severity Tiers"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = RGBColor(245, 158, 11)

tiers = [
    ("HIGH SEVERITY (Red Alert)", "Direct factual contradiction: renamed CLI flag, modified default value, altered function arguments. Blocks PR merge (Exit Code 1).", ACCENT_ROSE),
    ("MEDIUM SEVERITY (Warning)", "Indirect semantic shift or behavioral nuance change. Recommended review.", RGBColor(245, 158, 11)),
    ("LOW SEVERITY (Info)", "Minor wording or contextual reference update.", ACCENT_BLUE),
    ("Same-PR Fix Suppression", "If a developer modifies DEFAULT_TIMEOUT = 60 and updates the doc sentence to '60 seconds' in the same PR, DriftLens detects it and suppresses the finding.", ACCENT_GREEN)
]
for title, desc, col in tiers:
    p = rtf.add_paragraph()
    p.text = f"\n{title}"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = col
    p = rtf.add_paragraph()
    p.text = desc
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(2)


# ==========================================
# SLIDE 6: Golden Evaluation Demos
# ==========================================
s6 = create_base_slide("6 Golden Evaluation Demos (100% Problem Coverage)")

demos = [
    ("1. Renamed CLI Flag", "cli.py: --timeout -> --deadline\nDoc: 'Use --timeout to stop jobs'\nResult: FAIL (High, 98% Conf)\nFix: 'Use --deadline to stop jobs'"),
    ("2. Changed Default Value", "config.js: DEFAULT_TIMEOUT = 30 -> 60\nDoc: 'Default timeout is 30s'\nResult: FAIL (High, 93% Conf)\nFix: 'Default timeout is 60s'"),
    ("3. False-Positive Control", "utils.js: Internal math refactor\nDoc: 'Fast calculation utilities'\nResult: PASS (0 Findings)\nProves zero false alarms!"),
    ("4. Config Table Drift", "config.js: TIMEOUT_SECONDS = 30 -> 60\nDoc: | TIMEOUT_SECONDS | 30 |\nResult: FAIL (High, 97% Conf)\nFix: Updates table row to 60"),
    ("5. API Parameter Signature", "api.js: registerUser(email, password, tenantId)\nDoc: 'registerUser(email, password)'\nResult: FAIL (High, 91% Conf)\nFix: Updates function signature"),
    ("6. Deleted Feature / Route", "routes.js: Deleted /api/v1/legacy-export\nDoc: 'Use /api/v1/legacy-export for CSV'\nResult: FAIL (High, 90% Conf)\nFix: Flags obsolete documentation")
]

for idx, (title, desc) in enumerate(demos):
    row = idx // 3
    col_idx = idx % 3
    x = 0.8 + col_idx * 3.9
    y = 1.8 + row * 2.5
    card = add_card(s6, x, y, 3.8, 2.3)
    ctf = card.text_frame
    ctf.word_wrap = True
    p = ctf.paragraphs[0]
    p.text = title
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE if "3" not in title else ACCENT_GREEN
    p = ctf.add_paragraph()
    p.text = desc
    p.font.name = "Consolas"
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(6)


# ==========================================
# SLIDE 7: Standout Innovations
# ==========================================
s7 = create_base_slide("Standout Innovations: Beyond Drift Detection")

innovations = [
    ("✨ Starter README Generator", "What if a repo has zero documentation? DriftLens statically inspects all constants, CLI flags, and routes to synthesize an accurate starter README.md automatically.\n\nCLI: driftlens generate-readme --repo .", ACCENT_BLUE),
    ("🌐 Visual Blast Radius Graph", "Interactive SVG dependency graph maps code changes to affected documentation sections. Visualizes exactly where code changes radiate impact into documentation.", ACCENT_GREEN),
    ("⚡ 1-Click Auto-Remediation", "Maintainers don't just get an alert — they can click 'Apply Fix' directly in the Web Studio or download a git .patch file to apply the correction in 1 second.", RGBColor(245, 158, 11))
]

for idx, (title, desc, col) in enumerate(innovations):
    x = 0.8 + idx * 3.9
    card = add_card(s7, x, 1.8, 3.8, 4.8)
    ctf = card.text_frame
    ctf.word_wrap = True
    p = ctf.paragraphs[0]
    p.text = title
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = col
    p = ctf.add_paragraph()
    p.text = desc
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(14)


# ==========================================
# SLIDE 8: Verification & Production Readiness
# ==========================================
s8 = create_base_slide("Verification, Performance & CI/CD Integration")

stat_boxes = [
    ("70", "Unit & Integration Tests", ACCENT_GREEN),
    ("100%", "Test Pass Rate (Vitest)", ACCENT_BLUE),
    ("5ms", "Live Benchmark Duration", ACCENT_GREEN),
    ("Qwen 2.5 7B", "Open-Weight Model Used", RGBColor(245, 158, 11))
]

for idx, (big, label, col) in enumerate(stat_boxes):
    x = 0.8 + idx * 2.95
    card = add_card(s8, x, 1.8, 2.85, 1.6)
    ctf = card.text_frame
    ctf.word_wrap = True
    p = ctf.paragraphs[0]
    p.text = big
    p.font.name = "Arial"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = col
    p.alignment = PP_ALIGN.CENTER
    p = ctf.add_paragraph()
    p.text = label
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.alignment = PP_ALIGN.CENTER

# Bottom Two Cards
b_left = add_card(s8, 0.8, 3.7, 5.7, 3.0)
btf1 = b_left.text_frame
btf1.word_wrap = True
p = btf1.paragraphs[0]
p.text = "Reusable GitHub Action (action.yml)"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = btf1.add_paragraph()
p.text = "• Reusable GitHub Action for any open-source repo\n• Automates PR comments with diff suggestions\n• Fully supports GitHub Actions runners with Node 20+\n• Blocks PR merging with exit code 1 on HIGH severity drift"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)

b_right = add_card(s8, 6.8, 3.7, 5.7, 3.0)
btf2 = b_right.text_frame
btf2.word_wrap = True
p = btf2.paragraphs[0]
p.text = "Engineering Optimizations"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = btf2.add_paragraph()
p.text = "• O(1) Token Lookups: Pre-computed token and stem sets\n• Worker Pool Concurrency: Parallelized LLM inference\n• Parallel Async I/O: Fast repository traversal\n• In-Memory RAM Caching: Sub-millisecond API response"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)


# ==========================================
# SLIDE 9: Conclusion & Summary
# ==========================================
s9 = create_base_slide("Summary: Why DriftLens Wins Problem 5")

c_summary = add_card(s9, 0.8, 1.8, 11.7, 4.8)
stf = c_summary.text_frame
stf.word_wrap = True

p = stf.paragraphs[0]
p.text = "Key Takeaways for Judges"
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

takeaways = [
    ("1. Solves the Real Challenge", "Identifies documentation drift in explanatory prose, configuration tables, and API tutorials where standard code-block test runners are blind."),
    ("2. Pre-Correction Detection", "Flags the exact contradictory documentation sentence on the code commit, before any corrective documentation commit is required."),
    ("3. Zero Hallucinations", "Deterministic AST and diff parsing anchors all file paths and line ranges; the AI model only performs bounded semantic reasoning."),
    ("4. 100% Hackathon Rule Compliant", "Powered exclusively by open-weight Qwen 2.5 7B. Zero closed/proprietary APIs."),
    ("5. Complete Developer Experience", "CLI tool + Reusable GitHub Action + Interactive Web Studio with Visual Blast Radius Graph + 1-Click Remediation + Starter README Generator.")
]

for title, desc in takeaways:
    p = stf.add_paragraph()
    p.text = f"• {title}: {desc}"
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(10)

output_path = "/Users/yashvardhanmundhra/Desktop/driftlens/DriftLens_Presentation.pptx"
prs.save(output_path)
print(f"Presentation successfully generated at: {output_path}")
